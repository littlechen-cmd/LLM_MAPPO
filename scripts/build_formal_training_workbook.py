"""Render and verify the owner's Chinese core-experiment workbook."""

from __future__ import annotations

from hashlib import sha256
import json
import math
from pathlib import Path
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import yaml


ROOT = Path(__file__).resolve().parents[1]
CONTENT = ROOT / "docs/handoff/formal_training_workbook_content.json"
OUTPUT = ROOT / "docs/handoff/三场景核心实验训练配置.xlsx"
SEEDS = set(range(7, 78, 10))
SLUGS = {
    "baseline": "baseline", "A* KD": "astar-kd", "LLM KD": "llm-kd",
    "A* KD + LLM KD": "astar-llm-kd",
}


def make_sheet(book, item):
    sheet = book.create_sheet(item["title"])
    sheet.append(item["headers"])
    for row in item["rows"]:
        sheet.append(row)
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.tabColor = "17365D"
    sheet.row_dimensions[1].height = 34
    for index, width in enumerate(item["widths"], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        height = 30
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="253746")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EEF3F8")
            width = item["widths"][cell.column - 1]
            lines = math.ceil(len(str(cell.value)) / max(10, width / 1.5))
            height = max(height, 15 * min(12, lines))
            if isinstance(cell.value, int) and abs(cell.value) >= 1000:
                cell.number_format = "#,##0"
            if isinstance(cell.value, str) and cell.value.startswith("https://"):
                cell.hyperlink = cell.value
                cell.font = Font(name="Microsoft YaHei", size=10, color="0563C1")
        sheet.row_dimensions[row[0].row].height = height
    table = Table(displayName=f"ConfigTable{len(book.worksheets)}", ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
    sheet.add_table(table)


def build_workbook():
    content = json.loads(CONTENT.read_text(encoding="utf-8"))
    assert content["schema"] == "formal-training-workbook-content-v1"
    book = Workbook()
    book.remove(book.active)
    book.properties.title = content["version"]
    book.properties.subject = "三场景、四方法、八配对种子的配置与证据"
    book.properties.creator = "Codex / 项目核心架构师"
    book.properties.description = "参数决策已给定；技术依赖据实列明，未运行长训练。"
    for item in content["sheets"]:
        make_sheet(book, item)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUTPUT)
    verify_workbook(OUTPUT)
    return OUTPUT


def verify_workbook(path):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    book = load_workbook(path, read_only=True, data_only=True)
    assert len(book.sheetnames) == 16
    matrix = list(book["01_十二组配置"].values)[1:]
    runs = list(book["11_全部运行清单"].values)[1:]
    assert len(matrix) == 12 and len(runs) == 96
    assert {row[3] for row in runs} == SEEDS
    assert sum(row[4] for row in runs) == 100800000
    assert len({(row[1], row[2], row[3]) for row in runs}) == 96
    instances = list(book["08_seed实例清单"].values)[1:]
    evaluation = [row[2] for row in instances if row[0] == "主评估"]
    durability = [row[2] for row in instances if row[0] == "formal耐久"]
    assert len(evaluation) == len(set(evaluation)) == 200
    assert len(durability) == 10 and not set(evaluation) & set(durability)
    for row in matrix:
        template = (ROOT / "configs/formal_environment_training"
                    / f"{row[0]}-{SLUGS[row[1]]}.yaml")
        values = yaml.safe_load(template.read_bytes())
        env, training = values["environment"], values["training"]
        assert row[2:7] == (
            env["n_agents"], len(env["charging_stations"]),
            len(env["picking_stations"]), env["initial_active_tasks"], env["max_steps"],
        )
        assert row[7] == training["environment_steps"]
        assert row[8:13] == (16, 64, 1024, 8, 200)
        assert training["num_env_workers"] == 16 and training["rollout_length"] == 64
        assert row[14:16] == tuple(training["milestone_checkpoint_steps"])
        assert row[16] == values["method"]
    for row in list(book["12_来源身份"].values)[1:]:
        if row[0] in {"原始配置模板", "新短测证据快照"}:
            assert sha256((ROOT / row[1]).read_bytes()).hexdigest() == row[2]
    book.close()


if __name__ == "__main__":
    print(build_workbook())
